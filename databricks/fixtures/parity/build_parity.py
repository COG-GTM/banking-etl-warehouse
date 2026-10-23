"""Rebuild the legacy SQL Server DWH parity fixtures (Ticket 10).

Restores ``data_sources/sample.bak`` into a local SQL Server 2022 container, creates the
legacy ``DWH`` star schema with ``sql_scripts/01_create_tables.sql`` and loads it with
T-SQL that reproduces the four Talend jobs (``talend_jobs/*.zip`` ``.item`` XML), then
exports the four DWH tables and the two stored-procedure result sets to CSV.

Talend behaviour reproduced (from the .item XML, not the README):

* Load_DimBranch / Load_DimAccount: straight copies, ``INSERT``, DIE_ON_ERROR=false.
* Load_DimCustomer: tMap ``customer`` main + ``city`` and ``state`` lookups without
  ``innerJoin`` -> LEFT OUTER joins; ``StringHandling.UPCASE`` on name, address, gender.
* Load_FactTransaction: tUnite(SQL ``transaction_db``, Excel, CSV in mergeOrder 1/2/3)
  -> tUniqRow on transaction_id (first occurrence wins) -> TRUNCATE + row-level INSERT
  with DIE_ON_ERROR=false, so rows violating FK_FactTransaction_DimAccount /
  FK_FactTransaction_DimBranch are rejected and the job carries on.

The Talend loads run in the ``sample`` DB context and write to ``DWH.dbo.*``. Excel and
CSV rows are read from the committed files here and inserted one statement at a time
(sqlcmd continues after Msg 547), which is the legacy row-level reject behaviour.

    python databricks/fixtures/parity/build_parity.py            # uses container mssql-t10
    python databricks/fixtures/parity/build_parity.py --container my-mssql --password ...

Needs: docker, the mcr.microsoft.com/mssql/server:2022-latest image, openpyxl.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import subprocess
from datetime import datetime
from decimal import Decimal
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
DATA = REPO / "data_sources"
SQL = REPO / "sql_scripts"

# Legacy object -> (fixture file, SELECT producing the legacy result set in legacy column names)
EXPORTS = {
    "DimBranch": ("dim_branch.csv", "SELECT BranchID, BranchName, BranchLocation FROM DWH.dbo.DimBranch"),
    "DimAccount": ("dim_account.csv",
                   "SELECT AccountID, CustomerID, AccountType, Balance, DateOpened, Status FROM DWH.dbo.DimAccount"),
    "DimCustomer": ("dim_customer.csv",
                    "SELECT CustomerID, CustomerName, Address, CityName, StateName, Age, Gender, Email "
                    "FROM DWH.dbo.DimCustomer"),
    "FactTransaction": ("fact_transaction.csv",
                        "SELECT TransactionID, AccountID, TransactionDate, Amount, TransactionType, BranchID "
                        "FROM DWH.dbo.FactTransaction"),
}
PROC_EXPORTS = {
    # file -> (procedure call, temp table DDL)
    "sp_daily_transaction__all.csv": (
        "EXEC DWH.dbo.sp_DailyTransaction @start_date='0001-01-01', @end_date='9999-12-31'",
        "[Date] DATE, TotalTransactions INT, TotalAmount MONEY",
    ),
    "sp_daily_transaction__2024-01-18__2024-01-20.csv": (
        "EXEC DWH.dbo.sp_DailyTransaction @start_date='2024-01-18', @end_date='2024-01-20'",
        "[Date] DATE, TotalTransactions INT, TotalAmount MONEY",
    ),
    "sp_balance_per_customer__all.csv": (
        "EXEC DWH.dbo.sp_BalancePerCustomer @customer_name=''",
        "CustomerName VARCHAR(100), AccountType VARCHAR(50), InitialBalance MONEY, CurrentBalance MONEY",
    ),
    "sp_balance_per_customer__shelly.csv": (
        "EXEC DWH.dbo.sp_BalancePerCustomer @customer_name='Shelly'",
        "CustomerName VARCHAR(100), AccountType VARCHAR(50), InitialBalance MONEY, CurrentBalance MONEY",
    ),
}


class Sql:
    def __init__(self, container: str, password: str):
        self.container, self.password = container, password

    def run(self, script: str, db: str = "master", check: bool = True) -> str:
        cmd = ["docker", "exec", "-i", self.container, "/opt/mssql-tools18/bin/sqlcmd", "-C", "-S", "localhost",
               "-U", "sa", "-P", self.password, "-d", db, "-y", "0"]
        p = subprocess.run(cmd, input=script, capture_output=True, text=True)
        if check and p.returncode != 0:
            raise RuntimeError(p.stdout + p.stderr)
        return p.stdout

    def json_rows(self, select: str, db: str = "sample") -> list[dict]:
        *setup, query = [s for s in select.split("; ") if s]
        prefix = "; ".join(setup) + ";" if setup else ""
        out = self.run(f"SET NOCOUNT ON; {prefix} SELECT ISNULL(({query} FOR JSON PATH, INCLUDE_NULL_VALUES), '[]')",
                       db=db)
        text = next(line.strip() for line in out.splitlines() if line.strip().startswith("["))
        return json.loads(text, parse_float=Decimal)


def sql_literal(v) -> str:
    if v is None:
        return "NULL"
    if isinstance(v, datetime):
        return f"'{v:%Y-%m-%d %H:%M:%S}'"
    if isinstance(v, (int, Decimal)):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


def excel_rows() -> list[tuple]:
    """tFileInputExcel_1: Sheet1, header 1 row; date cells come back as java.util.Date."""
    from openpyxl import load_workbook

    ws = load_workbook(DATA / "transaction_excel.xlsx", data_only=True)["Sheet1"]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if all(c is None for c in r):
            continue
        tid, acc, ts, amt, ttype, br = r[:6]
        if isinstance(ts, datetime) and ts.microsecond:  # Excel serials carry float noise; Talend rounds to ms
            ts = datetime.fromtimestamp(round(ts.timestamp(), 3))
        rows.append((int(tid), int(acc), ts, int(amt), ttype, int(br)))
    return rows


def csv_rows() -> list[tuple]:
    """tFileInputDelimited_1: ',' separator, header 1, transaction_date pattern dd-MM-yyyy HH:mm:ss."""
    rows = []
    with open(DATA / "transaction_csv.csv", encoding="ISO-8859-1", newline="") as fh:
        for r in csv.DictReader(fh):
            rows.append((int(r["transaction_id"]), int(r["account_id"]),
                         datetime.strptime(r["transaction_date"], "%d-%m-%Y %H:%M:%S"),
                         int(r["amount"]), r["transaction_type"], int(r["branch_id"])))
    return rows


def talend_equivalent_load(sql: Sql) -> dict:
    sql.run(f"USE sample;\n{(SQL / '01_create_tables.sql').read_text(encoding='utf-8')}")
    procs = (SQL / "02_create_procedures.sql").read_text(encoding="utf-8").replace(
        "CREATE PROCEDURE", "GO\nCREATE PROCEDURE")  # the committed script lacks a batch separator (Msg 111)
    sql.run(procs)

    # Load_DimBranch, Load_DimAccount, Load_DimCustomer (sample DB context, plain INSERT).
    sql.run("""
INSERT INTO DWH.dbo.DimBranch (BranchID, BranchName, BranchLocation)
SELECT branch_id, branch_name, branch_location FROM dbo.branch;
INSERT INTO DWH.dbo.DimAccount (AccountID, CustomerID, AccountType, Balance, DateOpened, Status)
SELECT account_id, customer_id, account_type, balance, date_opened, status FROM dbo.account;
INSERT INTO DWH.dbo.DimCustomer (CustomerID, CustomerName, Address, CityName, StateName, Age, Gender, Email)
SELECT c.customer_id, UPPER(c.customer_name), UPPER(c.address), ci.city_name, s.state_name,
       c.age, UPPER(c.gender), c.email
FROM dbo.customer c
LEFT JOIN dbo.city ci ON ci.city_id = c.city_id
LEFT JOIN dbo.state s ON s.state_id = ci.state_id;
""", db="sample")

    # Load_FactTransaction: tUnite (SQL, Excel, CSV) -> tUniqRow(transaction_id) -> TRUNCATE + INSERT.
    sql_stream = [(r["transaction_id"], r["account_id"], datetime.fromisoformat(r["transaction_date"]),
                   r["amount"], r["transaction_type"], r["branch_id"])
                  for r in sql.json_rows("SELECT transaction_id, account_id, transaction_date, amount, "
                                         "transaction_type, branch_id FROM dbo.transaction_db")]
    streams = [("sqlserver", sql_stream), ("excel", excel_rows()), ("csv", csv_rows())]
    seen, unique, dropped = set(), [], []
    for name, rows in streams:
        for row in rows:
            (dropped if row[0] in seen else unique).append((name, row))
            seen.add(row[0])
    stmts = ["TRUNCATE TABLE DWH.dbo.FactTransaction;"]
    for _, row in unique:
        stmts.append("INSERT INTO DWH.dbo.FactTransaction (TransactionID, AccountID, TransactionDate, Amount, "
                     f"TransactionType, BranchID) VALUES ({', '.join(sql_literal(v) for v in row)});\nGO")
    out = sql.run("\n".join(stmts), db="sample", check=False)
    loaded = {r["TransactionID"] for r in sql.json_rows("SELECT TransactionID FROM DWH.dbo.FactTransaction")}
    rejected = [{"transaction_id": row[0], "account_id": row[1], "branch_id": row[5], "stream": name}
                for name, row in unique if row[0] not in loaded]
    return {
        "stream_rows": {name: len(rows) for name, rows in streams},
        "after_tUniqRow": len(unique),
        "duplicates_dropped": [{"transaction_id": row[0], "stream": name} for name, row in dropped],
        "insert_rejected_fk": rejected,
        "insert_errors": sorted({line.split(",")[0] for line in out.splitlines() if line.startswith("Msg ")}),
    }


def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, Decimal):
        return f"{v:.4f}"
    if isinstance(v, str) and len(v) >= 19 and v[10] == "T":
        return v.replace("T", " ")
    return str(v)


def write_csv(path: Path, columns: list[str], rows: list[dict]) -> None:
    with open(path, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(columns)
        for r in rows:
            w.writerow([_fmt(r.get(c)) for c in columns])


def export(sql: Sql, load_report: dict) -> dict:
    manifest = {"generated_by": "databricks/fixtures/parity/build_parity.py", "null_encoding": "empty field",
                "sql_server": sql.run("SET NOCOUNT ON; SELECT @@VERSION").strip().splitlines()[0],
                "inputs": {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                           for p in (DATA / "sample.bak", DATA / "transaction_excel.xlsx",
                                     DATA / "transaction_csv.csv")},
                "talend_fact_load": load_report, "files": {}}
    for obj, (fname, select) in EXPORTS.items():
        rows = sql.json_rows(select + " ORDER BY 1")
        cols = [c.strip() for c in select.split("SELECT ", 1)[1].split(" FROM ")[0].split(",")]
        write_csv(HERE / fname, cols, rows)
        manifest["files"][fname] = {"legacy_object": f"DWH.dbo.{obj}", "rows": len(rows)}
    for fname, (call, ddl) in PROC_EXPORTS.items():
        cols = [c.split()[0].strip("[]") for c in ddl.split(",")]
        rows = sql.json_rows(f"CREATE TABLE #r ({ddl}); INSERT INTO #r {call}; "
                             f"SELECT {', '.join('[' + c + ']' for c in cols)} FROM #r ORDER BY 1, 2")
        write_csv(HERE / fname, cols, rows)
        manifest["files"][fname] = {"legacy_object": call, "rows": len(rows)}
    (HERE / "manifest.json").write_text(json.dumps(manifest, indent=2, default=str) + "\n", encoding="utf-8")
    return manifest


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--container", default="mssql-t10")
    p.add_argument("--password", default="T10_Strong!Pass")
    p.add_argument("--skip-restore", action="store_true")
    a = p.parse_args()
    sql = Sql(a.container, a.password)
    if not a.skip_restore:
        subprocess.run(["docker", "cp", str(DATA / "sample.bak"), f"{a.container}:/tmp/sample.bak"], check=True)
        sql.run("IF DB_ID('DWH') IS NOT NULL BEGIN ALTER DATABASE DWH SET SINGLE_USER WITH ROLLBACK IMMEDIATE; "
                "DROP DATABASE DWH; END\nGO\n"
                "IF DB_ID('sample') IS NOT NULL BEGIN ALTER DATABASE sample SET SINGLE_USER WITH ROLLBACK IMMEDIATE; "
                "DROP DATABASE sample; END\nGO\n"
                "RESTORE DATABASE sample FROM DISK='/tmp/sample.bak' WITH "
                "MOVE 'sample' TO '/var/opt/mssql/data/sample.mdf', "
                "MOVE 'sample_log' TO '/var/opt/mssql/data/sample_log.ldf'")
    report = talend_equivalent_load(sql)
    manifest = export(sql, report)
    print(json.dumps({k: v["rows"] for k, v in manifest["files"].items()}, indent=2))
    print(json.dumps(report, indent=2, default=str))


if __name__ == "__main__":
    main()
