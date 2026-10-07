from __future__ import annotations

import itertools
import json
import shutil
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from banking_etl.bronze import files
from banking_etl.bronze.files import (
    CSV_SOURCE,
    EXCEL_SOURCE,
    INGESTED_AT_COLUMN,
    RESCUED_DATA_COLUMN,
    SOURCE_FILE_COLUMN,
    TRANSACTION_SCHEMA,
    Locations,
)

_counter = itertools.count()


@pytest.fixture
def loc(spark, tmp_path) -> Locations:
    prefix = f"t4_{next(_counter)}_"
    loc = Locations(catalog=None, schema_prefix=prefix, landing_root=str(tmp_path / "landing"))
    spark.sql(f"CREATE DATABASE IF NOT EXISTS {loc.bronze_schema}")
    return loc


def stage_csv(loc: Locations, src: Path, name: str | None = None) -> Path:
    folder = Path(loc.source_path(CSV_SOURCE))
    folder.mkdir(parents=True, exist_ok=True)
    return Path(shutil.copy(src, folder / (name or src.name)))


def excel_rows(path: Path) -> list[tuple]:
    ws = openpyxl.load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))[1:]
    return [(int(a), int(b), c, Decimal(d), e, int(f)) for a, b, c, d, e, f in rows]


def as_excel_text(value) -> str | None:
    """Cell text as the Databricks Excel reader returns it for a STRING hint (e.g. 1/18/24 13:10)."""
    if isinstance(value, datetime):
        return f"{value.month}/{value.day}/{value:%y} {value.hour}:{value:%M}"
    return None if value is None else str(value)


def stage_excel(spark, loc: Locations, xlsx: Path, name: str = "part-0") -> None:
    ws = openpyxl.load_workbook(xlsx, read_only=True).active
    rows = [tuple(as_excel_text(v) for v in r) for r in list(ws.iter_rows(values_only=True))[1:]]
    stage_excel_rows(spark, loc, rows, name)


def stage_excel_rows(spark, loc: Locations, rows: list[tuple], name: str = "part-0") -> None:
    """All-STRING parquet stand-in for the workbook: OSS Spark has no Excel reader."""
    out = Path(loc.source_path(EXCEL_SOURCE)) / name
    schema = ", ".join(f"{c} STRING" for c in files.TRANSACTION_COLUMNS)
    spark.createDataFrame(rows, schema).coalesce(1).write.parquet(str(out))
    # The file source treats each parquet file as a landed file; flatten into the folder.
    for f in out.glob("*.parquet"):
        f.rename(out.parent / f"{name}.parquet")
    shutil.rmtree(out)


def run(spark, loc: Locations, names: list[str]) -> dict[str, int]:
    return files.ingest(spark, loc, names, reader=files.read_local_standin)


def test_schema_hints_cover_legacy_fact_columns():
    assert files.schema_hints() == (
        "transaction_id INT, account_id INT, transaction_date TIMESTAMP, "
        "amount DECIMAL(19,4), transaction_type STRING, branch_id INT"
    )


def test_default_locations_follow_shared_conventions():
    loc = Locations()
    assert loc.table_name(CSV_SOURCE) == "migration_demo.banking_mig_bronze.file_transaction_csv"
    assert loc.table_name(EXCEL_SOURCE) == "migration_demo.banking_mig_bronze.file_transaction_excel"
    landing = "/Volumes/migration_demo/banking_mig_bronze/landing"
    assert loc.source_path(CSV_SOURCE) == f"{landing}/transaction_csv"
    assert loc.schema_location(EXCEL_SOURCE) == f"{landing}/_autoloader/schemas/file_transaction_excel"
    assert loc.checkpoint_location(CSV_SOURCE) == f"{landing}/_autoloader/checkpoints/file_transaction_csv"
    assert not loc.checkpoint_location(CSV_SOURCE).startswith(loc.source_path(CSV_SOURCE) + "/")


def test_autoloader_options_csv():
    opts = files.autoloader_options(CSV_SOURCE, "/x/schema")
    assert opts["cloudFiles.format"] == "csv"
    assert opts["cloudFiles.schemaLocation"] == "/x/schema"
    assert opts["cloudFiles.schemaHints"] == files.schema_hints()
    assert opts["cloudFiles.schemaEvolutionMode"] == "rescue"
    assert opts["rescuedDataColumn"] == RESCUED_DATA_COLUMN
    assert opts["timestampFormat"] == "dd-MM-yyyy HH:mm:ss"
    assert opts["header"] == "true"
    assert opts["encoding"] == "ISO-8859-1"


def test_autoloader_options_excel():
    opts = files.autoloader_options(EXCEL_SOURCE, "/x/schema")
    assert opts["cloudFiles.format"] == "excel"
    assert opts["headerRows"] == "1"
    assert opts["dataAddress"] == "Sheet1"
    assert opts["cloudFiles.schemaEvolutionMode"] == "none"
    assert opts["rescuedDataColumn"] == RESCUED_DATA_COLUMN
    # Typed hints make the Excel reader fail the file on a bad cell; cast in to_bronze instead.
    assert opts["cloudFiles.schemaHints"] == files.schema_hints(as_strings=True)
    assert opts["cloudFiles.inferColumnTypes"] == "false"


def test_bronze_table_schema(spark, loc, data_sources):
    stage_csv(loc, data_sources / "transaction_csv.csv")
    run(spark, loc, ["transaction_csv"])
    schema = spark.table(loc.table_name(CSV_SOURCE)).schema
    assert [f.name for f in schema.fields] == [
        *files.TRANSACTION_COLUMNS,
        RESCUED_DATA_COLUMN,
        SOURCE_FILE_COLUMN,
        INGESTED_AT_COLUMN,
    ]
    types = {f.name: f.dataType.simpleString() for f in schema.fields}
    assert types["transaction_date"] == "timestamp"
    assert types["amount"] == "decimal(19,4)"
    assert types[INGESTED_AT_COLUMN] == "timestamp"


def test_csv_ingest_parses_dd_mm_yyyy_and_stamps_metadata(spark, loc, data_sources):
    stage_csv(loc, data_sources / "transaction_csv.csv")
    counts = run(spark, loc, ["transaction_csv"])
    assert counts == {loc.table_name(CSV_SOURCE): 12}

    df = spark.table(loc.table_name(CSV_SOURCE))
    rows = {r.transaction_id: r for r in df.collect()}
    assert sorted(rows) == list(range(14, 26))
    first = rows[14]
    assert first.transaction_date == datetime(2024, 1, 21, 14, 0)
    assert first.amount == Decimal("1500000")
    assert (first.account_id, first.transaction_type, first.branch_id) == (13, "Deposit", 4)
    assert all(r[RESCUED_DATA_COLUMN] is None for r in rows.values())
    assert all(r[SOURCE_FILE_COLUMN].endswith("/transaction_csv/transaction_csv.csv") for r in rows.values())
    assert all(r[INGESTED_AT_COLUMN] is not None for r in rows.values())


def test_excel_ingest_matches_workbook(spark, loc, data_sources):
    xlsx = data_sources / "transaction_excel.xlsx"
    stage_excel(spark, loc, xlsx)
    counts = run(spark, loc, ["transaction_excel"])
    assert counts == {loc.table_name(EXCEL_SOURCE): 7}

    got = sorted(
        tuple(r[c] for c in files.TRANSACTION_COLUMNS)
        for r in spark.table(loc.table_name(EXCEL_SOURCE)).collect()
    )
    assert got == sorted(excel_rows(xlsx))
    meta = spark.table(loc.table_name(EXCEL_SOURCE)).select(SOURCE_FILE_COLUMN).distinct().collect()
    assert [m[0].rsplit("/", 1)[-1] for m in meta] == ["part-0.parquet"]


def test_both_sources_in_one_run(spark, loc, data_sources):
    stage_csv(loc, data_sources / "transaction_csv.csv")
    stage_excel(spark, loc, data_sources / "transaction_excel.xlsx")
    counts = run(spark, loc, list(files.SOURCES))
    assert counts == {loc.table_name(EXCEL_SOURCE): 7, loc.table_name(CSV_SOURCE): 12}


def test_rerun_is_idempotent_and_new_files_are_incremental(spark, loc, data_sources, tmp_path):
    stage_csv(loc, data_sources / "transaction_csv.csv")
    assert run(spark, loc, ["transaction_csv"])[loc.table_name(CSV_SOURCE)] == 12
    assert run(spark, loc, ["transaction_csv"])[loc.table_name(CSV_SOURCE)] == 12

    extra = tmp_path / "extra.csv"
    extra.write_text(
        "transaction_id,account_id,transaction_date,amount,transaction_type,branch_id\n"
        "26,24,23-01-2024 09:15:00,250000,Deposit,2\n"
    )
    stage_csv(loc, extra, "transaction_csv_20240123.csv")
    assert run(spark, loc, ["transaction_csv"])[loc.table_name(CSV_SOURCE)] == 13
    new = spark.table(loc.table_name(CSV_SOURCE)).where("transaction_id = 26").collect()
    assert len(new) == 1
    assert new[0][SOURCE_FILE_COLUMN].endswith("transaction_csv_20240123.csv")
    assert new[0].transaction_date == datetime(2024, 1, 23, 9, 15)


def test_malformed_values_land_in_rescued_data(spark, loc, tmp_path):
    bad = tmp_path / "bad.csv"
    bad.write_text(
        "transaction_id,account_id,transaction_date,amount,transaction_type,branch_id\n"
        "30,1,2024/01/22 09:00,100,Deposit,1\n"
        "31,1,22-01-2024 09:00:00,100,Deposit,1\n"
    )
    stage_csv(loc, bad)
    run(spark, loc, ["transaction_csv"])
    rows = {r.transaction_id: r for r in spark.table(loc.table_name(CSV_SOURCE)).collect()}
    assert rows[31][RESCUED_DATA_COLUMN] is None
    assert rows[31].transaction_date == datetime(2024, 1, 22, 9, 0)
    assert rows[30].transaction_date is None
    assert "2024/01/22 09:00" in rows[30][RESCUED_DATA_COLUMN]


def test_custom_landing_root_isolates_autoloader_state_per_target():
    a = Locations(catalog="cat_a", landing_root="/Volumes/shared/landing")
    b = Locations(catalog="cat_b", landing_root="/Volumes/shared/landing")
    assert a.source_path(CSV_SOURCE) == b.source_path(CSV_SOURCE)
    assert a.checkpoint_location(CSV_SOURCE) != b.checkpoint_location(CSV_SOURCE)
    assert a.schema_location(CSV_SOURCE) != b.schema_location(CSV_SOURCE)
    assert a.checkpoint_location(CSV_SOURCE) == (
        "/Volumes/shared/landing/_autoloader/cat_a.banking_mig_bronze/checkpoints/file_transaction_csv"
    )


def test_shared_landing_root_fills_each_target(spark, tmp_path, data_sources):
    root = str(tmp_path / "landing")
    targets = [Locations(catalog=None, schema=f"t4_shared_{i}", landing_root=root) for i in (0, 1)]
    stage_csv(targets[0], data_sources / "transaction_csv.csv")
    for loc in targets:
        spark.sql(f"CREATE DATABASE IF NOT EXISTS {loc.bronze_schema}")
        assert run(spark, loc, ["transaction_csv"]) == {loc.table_name(CSV_SOURCE): 12}


def test_schema_override():
    loc = Locations(schema="banking_mig_t4")
    assert loc.table_name(EXCEL_SOURCE) == "migration_demo.banking_mig_t4.file_transaction_excel"
    assert loc.landing == "/Volumes/migration_demo/banking_mig_t4/landing"


def test_empty_source_list_ingests_nothing(spark, loc, data_sources):
    stage_csv(loc, data_sources / "transaction_csv.csv")
    assert run(spark, loc, []) == {}
    assert not spark.catalog.tableExists(loc.table_name(CSV_SOURCE))


@pytest.mark.parametrize(
    "kwargs",
    [
        {"catalog": "x; DROP SCHEMA y"},
        {"schema_prefix": "a`b_"},
        {"schema": "s.t"},
    ],
)
def test_identifiers_are_validated(kwargs):
    with pytest.raises(ValueError):
        Locations(**kwargs)


def test_excel_malformed_cells_land_in_rescued_data(spark, loc):
    stage_excel_rows(
        spark,
        loc,
        [
            ("40", "1", "2024-01-22 9:00:00", "100", "Deposit", "1"),
            ("41", "1", "not a date", "200", "Deposit", "1"),
            ("42", "x", "1/22/24 10:00", "lots", "Payment", "2"),
        ],
    )
    assert run(spark, loc, ["transaction_excel"]) == {loc.table_name(EXCEL_SOURCE): 3}
    rows = {r.transaction_id: r for r in spark.table(loc.table_name(EXCEL_SOURCE)).collect()}
    assert rows[40].transaction_date == datetime(2024, 1, 22, 9, 0)
    assert rows[40][RESCUED_DATA_COLUMN] is None
    assert rows[41].transaction_date is None
    assert json.loads(rows[41][RESCUED_DATA_COLUMN]) == {
        "transaction_date": "not a date",
        "_file_path": rows[41][SOURCE_FILE_COLUMN],
    }
    assert (rows[42].account_id, rows[42].amount) == (None, None)
    assert rows[42].transaction_date == datetime(2024, 1, 22, 10, 0)
    rescued = json.loads(rows[42][RESCUED_DATA_COLUMN])
    assert (rescued["account_id"], rescued["amount"]) == ("x", "lots")
    assert "transaction_date" not in rescued
