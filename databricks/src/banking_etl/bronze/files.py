"""Incremental ingestion of the transaction CSV and Excel files into bronze (ticket 3).

Replaces the ``tFileInputDelimited`` and ``tFileInputExcel`` inputs of the Talend
``Load_FactTransaction`` job. On Databricks both files are read with Auto Loader
(``cloudFiles``) from the landing volume and appended to
``bronze.file_transaction_csv`` / ``bronze.file_transaction_excel`` with
``trigger(availableNow=True)``. The Auto Loader checkpoint guarantees each file is
ingested exactly once, so re-running is a no-op until a new file lands.

Locally (OSS Spark, no Auto Loader / Excel reader) ``read_transaction_csv`` and
``read_transaction_excel`` produce the same columns and metadata with batch reads,
so the parsing rules are unit-testable.
"""
from __future__ import annotations

import os
import shutil
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Callable, Iterable, Sequence

from pyspark.sql import DataFrame, SparkSession
from pyspark.sql import functions as F
from pyspark.sql.types import (
    DoubleType,
    LongType,
    StringType,
    StructField,
    StructType,
    TimestampNTZType,
)

from banking_etl.config import BRONZE_FILE_TRANSACTION_CSV, BRONZE_FILE_TRANSACTION_EXCEL, Settings

CSV = "csv"
EXCEL = "excel"
SOURCES = (CSV, EXCEL)

SOURCE_TAG = {CSV: "file_csv", EXCEL: "file_excel"}
TARGET_TABLE = {CSV: BRONZE_FILE_TRANSACTION_CSV, EXCEL: BRONZE_FILE_TRANSACTION_EXCEL}
# Landing folders under the bronze landing volume and the file names shipped in data_sources/.
LANDING_SUBDIR = {CSV: ("transactions", "csv"), EXCEL: ("transactions", "excel")}
SOURCE_FILE_NAME = {CSV: "transaction_csv.csv", EXCEL: "transaction_excel.xlsx"}

TRANSACTION_COLUMNS = ("transaction_id", "account_id", "transaction_date", "amount", "transaction_type", "branch_id")
METADATA_COLUMNS = ("_rescued_data", "_ingested_at", "_source", "_source_file")
RESCUED_DATA_COLUMN = "_rescued_data"

# The CSV uses dd-MM-yyyy HH:mm:ss, which inference would not parse; keep it raw in bronze
# (silver parses it with the Talend pattern).
CSV_SCHEMA_HINTS = "transaction_date STRING"
# Talend tFileInputDelimited: header 1, "," separator, "\n" rows, CSV_OPTION=false, ISO-8859-15.
# Spark 4's CSV reader rejects ISO-8859-15, so use ISO-8859-1 (identical except 8 code points such as the euro sign).
CSV_READER_OPTIONS = {"header": "true", "sep": ",", "encoding": "ISO-8859-1"}
# Talend tFileInputExcel: sheet "Sheet1", header 1, footer 0, first column 1.
EXCEL_SHEET = "Sheet1"
EXCEL_HEADER_ROWS = 1
EXCEL_READER_OPTIONS = {"headerRows": str(EXCEL_HEADER_ROWS), "dataAddress": EXCEL_SHEET}


@dataclass(frozen=True)
class FileSource:
    """Resolved locations for one file source."""

    name: str
    table: str
    landing_dir: str
    schema_location: str
    checkpoint_location: str


def file_source(settings: Settings, name: str) -> FileSource:
    if name not in SOURCES:
        raise ValueError(f"unknown file source {name!r}; expected one of {SOURCES}")
    table_name = TARGET_TABLE[name]
    return FileSource(
        name=name,
        table=settings.table("bronze", table_name),
        landing_dir=settings.landing_path(*LANDING_SUBDIR[name]),
        schema_location=settings.checkpoint_path("bronze", table_name, "_schema"),
        checkpoint_location=settings.checkpoint_path("bronze", table_name, "_checkpoint"),
    )


def parse_sources(value: str | Iterable[str] | None) -> tuple[str, ...]:
    """``"csv,excel"`` / ``["csv"]`` / empty -> validated tuple (empty means all)."""
    items = value.split(",") if isinstance(value, str) else list(value or [])
    names = tuple(dict.fromkeys(s.strip().lower() for s in items if s and s.strip()))
    for n in names:
        if n not in SOURCES:
            raise ValueError(f"unknown file source {n!r}; expected one of {SOURCES}")
    return names or SOURCES


# --------------------------------------------------------------------------- Auto Loader options


def csv_autoloader_options(src: FileSource) -> dict[str, str]:
    """Auto Loader options for the CSV: inference + evolution, raw transaction_date, rescued data."""
    return {
        "cloudFiles.format": "csv",
        "cloudFiles.schemaLocation": src.schema_location,
        "cloudFiles.inferColumnTypes": "true",
        "cloudFiles.schemaHints": CSV_SCHEMA_HINTS,
        "cloudFiles.schemaEvolutionMode": "addNewColumns",
        "cloudFiles.includeExistingFiles": "true",
        "rescuedDataColumn": RESCUED_DATA_COLUMN,
        **CSV_READER_OPTIONS,
    }


def excel_autoloader_options(src: FileSource) -> dict[str, str]:
    """Auto Loader options for native Excel (DBR 17.1+): schema evolution is unsupported -> ``none``."""
    return {
        "cloudFiles.format": "excel",
        "cloudFiles.schemaLocation": src.schema_location,
        "cloudFiles.inferColumnTypes": "true",
        "cloudFiles.schemaEvolutionMode": "none",
        "cloudFiles.includeExistingFiles": "true",
        **EXCEL_READER_OPTIONS,
    }


AUTOLOADER_OPTIONS: dict[str, Callable[[FileSource], dict[str, str]]] = {
    CSV: csv_autoloader_options,
    EXCEL: excel_autoloader_options,
}


# --------------------------------------------------------------------------- shared shaping


def with_file_metadata(df: DataFrame, source: str) -> DataFrame:
    """Append bronze metadata: ``_rescued_data`` (if absent), ``_ingested_at``, ``_source``, ``_source_file``.

    ``_source_file`` comes from the hidden ``_metadata.file_path`` column of file sources.
    """
    if RESCUED_DATA_COLUMN not in df.columns:
        df = df.withColumn(RESCUED_DATA_COLUMN, F.lit(None).cast(StringType()))
    data_cols = [c for c in df.columns if c not in METADATA_COLUMNS and c != "_metadata"]
    return df.select(
        *data_cols,
        F.col(RESCUED_DATA_COLUMN),
        F.current_timestamp().alias("_ingested_at"),
        F.lit(SOURCE_TAG[source]).alias("_source"),
        F.col("_metadata.file_path").alias("_source_file"),
    )


def apply_schema_hints(schema: StructType, hints: str) -> StructType:
    """Override inferred field types with ``"col TYPE, ..."`` hints (what ``cloudFiles.schemaHints`` does)."""
    hinted = {f.name.lower(): f for f in StructType.fromDDL(hints).fields}
    return StructType(
        [StructField(f.name, hinted[f.name.lower()].dataType, True) if f.name.lower() in hinted else f for f in schema]
    )


# --------------------------------------------------------------------------- batch readers (local / ad hoc)


def read_transaction_csv(spark: SparkSession, path: str) -> DataFrame:
    """Batch-read transaction CSV(s) with the same inference, hints and metadata as the Auto Loader stream."""
    reader = spark.read.options(**CSV_READER_OPTIONS)
    inferred = reader.option("inferSchema", "true").csv(path).schema
    df = reader.schema(apply_schema_hints(inferred, CSV_SCHEMA_HINTS)).csv(path)
    return with_file_metadata(df, CSV)


def parse_excel_bytes(
    content: bytes, *, sheet: str = EXCEL_SHEET, header_rows: int = EXCEL_HEADER_ROWS
) -> tuple[list[str], list[tuple]]:
    """Parse one sheet of an .xlsx into (header, rows) with openpyxl. Fully empty rows are skipped."""
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"sheet {sheet!r} not found; workbook has {wb.sheetnames}")
        rows = [r for r in wb[sheet].iter_rows(values_only=True) if any(v is not None for v in r)]
    finally:
        wb.close()
    if header_rows != 1:
        raise ValueError("only one header row is supported (same as the Databricks Excel reader)")
    if not rows:
        return [], []
    header = [str(h).strip() for h in rows[0]]
    width = len(header)
    return header, [tuple(r[:width]) + (None,) * (width - len(r)) for r in rows[1:]]


def _excel_type(values: Sequence[object]):
    present = [v for v in values if v is not None]
    if present and all(isinstance(v, bool) for v in present):
        return StringType()
    if present and all(isinstance(v, int) and not isinstance(v, bool) for v in present):
        return LongType()
    if present and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in present):
        return DoubleType()
    if present and all(isinstance(v, (datetime, date)) for v in present):
        return TimestampNTZType()
    return StringType()


def infer_excel_schema(header: Sequence[str], rows: Sequence[tuple]) -> StructType:
    """Mirror the Databricks Excel reader's inference: integers -> BIGINT, numbers -> DOUBLE,
    dates -> TIMESTAMP_NTZ, anything else -> STRING."""
    return StructType(
        [StructField(name, _excel_type([r[i] for r in rows]), True) for i, name in enumerate(header)]
    )


def _coerce(value, dtype):
    if value is None:
        return None
    if isinstance(dtype, StringType):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else str(value)
    if isinstance(dtype, DoubleType):
        return float(value)
    if isinstance(dtype, TimestampNTZType) and not isinstance(value, datetime):
        return datetime(value.year, value.month, value.day)
    return value


def read_transaction_excel(spark: SparkSession, path: str, *, sheet: str = EXCEL_SHEET) -> DataFrame:
    """Batch-read one .xlsx into the same shape as ``bronze.file_transaction_excel``.

    Uses the native Databricks Excel reader when present, else openpyxl on the driver
    (local tests; OSS Spark has no Excel data source).
    """
    try:
        df = spark.read.format("excel").options(**{**EXCEL_READER_OPTIONS, "dataAddress": sheet}).load(path)
        return with_file_metadata(df, EXCEL)
    except Exception as exc:  # noqa: BLE001 - OSS Spark: "Failed to find the data source: excel"
        if "excel" not in str(exc).lower():
            raise
    header, rows = parse_excel_bytes(Path(path).read_bytes(), sheet=sheet)
    schema = infer_excel_schema(header, rows)
    data = [tuple(_coerce(v, f.dataType) for v, f in zip(r, schema.fields)) for r in rows]
    df = spark.createDataFrame(data, schema).withColumn(
        "_metadata", F.struct(F.lit(Path(path).resolve().as_uri()).alias("file_path"))
    )
    return with_file_metadata(df, EXCEL)


# --------------------------------------------------------------------------- streaming (Databricks)


def read_autoloader(spark: SparkSession, src: FileSource) -> DataFrame:
    """Auto Loader stream over the source's landing folder, with bronze metadata columns."""
    df = spark.readStream.format("cloudFiles").options(**AUTOLOADER_OPTIONS[src.name](src)).load(src.landing_dir)
    return with_file_metadata(df, src.name)


def write_available_now(df: DataFrame, table: str, checkpoint_location: str, *, query_name: str | None = None):
    """Append a stream to a Delta table, processing everything available once, then stop.

    ``mergeSchema`` lets columns added by Auto Loader schema evolution reach the table.
    Returns the finished StreamingQuery.
    """
    writer = (
        df.writeStream.format("delta")
        .outputMode("append")
        .option("checkpointLocation", checkpoint_location)
        .option("mergeSchema", "true")
        .trigger(availableNow=True)
    )
    if query_name:
        writer = writer.queryName(query_name)
    query = writer.toTable(table)
    query.awaitTermination()
    return query


def _table_rows(spark: SparkSession, table: str) -> int:
    return spark.table(table).count() if spark.catalog.tableExists(table) else 0


def is_schema_change_error(exc: BaseException) -> bool:
    """Auto Loader stops the stream when ``addNewColumns`` records a new column; a restart picks it up."""
    text = f"{type(exc).__name__}: {exc}".upper()
    return "UNKNOWN_FIELD" in text or "UNKNOWNFIELDEXCEPTION" in text


def ingest_source(spark: SparkSession, src: FileSource, *, schema_change_retries: int = 1) -> dict:
    """Run one Auto Loader availableNow pass for ``src`` and report rows before/after.

    A schema-evolution stop (new column detected) is retried ``schema_change_retries`` times, so new
    columns land in the same run instead of failing the job.
    """
    before = _table_rows(spark, src.table)
    attempt = 0
    while True:
        try:
            write_available_now(read_autoloader(spark, src), src.table, src.checkpoint_location, query_name=f"bronze_{src.name}")
            break
        except Exception as exc:  # noqa: BLE001 - re-raised unless it is a schema-evolution restart
            if attempt >= schema_change_retries or not is_schema_change_error(exc):
                raise
            attempt += 1
    after = _table_rows(spark, src.table)
    return {
        "source": src.name,
        "table": src.table,
        "rows_before": before,
        "rows_after": after,
        "rows_added": after - before,
        "schema_change_restarts": attempt,
    }


def ingest_file_source(spark: SparkSession, settings: Settings, name: str, *, schema_change_retries: int = 1) -> dict:
    return ingest_source(spark, file_source(settings, name), schema_change_retries=schema_change_retries)


def ingest_transaction_files(
    spark: SparkSession,
    settings: Settings,
    sources: str | Iterable[str] | None = None,
    *,
    schema_change_retries: int = 1,
) -> list[dict]:
    return [
        ingest_file_source(spark, settings, name, schema_change_retries=schema_change_retries)
        for name in parse_sources(sources)
    ]


# --------------------------------------------------------------------------- landing helper


Copier = Callable[[str, str], None]
Exists = Callable[[str], bool]


def _local_copy(src: str, dst: str) -> None:
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    shutil.copyfile(src, dst)


def landing_targets(settings: Settings, source_dir: str | os.PathLike, sources=None) -> list[tuple[str, str]]:
    """(local file, landing volume path) for each requested source file in ``source_dir``."""
    return [
        (str(Path(source_dir) / SOURCE_FILE_NAME[n]), settings.landing_path(*LANDING_SUBDIR[n], SOURCE_FILE_NAME[n]))
        for n in parse_sources(sources)
    ]


def land_transaction_files(
    settings: Settings,
    source_dir: str | os.PathLike,
    *,
    sources=None,
    overwrite: bool = False,
    copy: Copier = _local_copy,
    exists: Exists = os.path.exists,
) -> list[tuple[str, str]]:
    """Copy ``transaction_csv.csv`` / ``transaction_excel.xlsx`` into the landing volume.

    Default copier uses the /Volumes FUSE mount (works inside Databricks); the laptop script passes a
    ``databricks fs cp`` copier. Existing targets are skipped unless ``overwrite`` (Auto Loader ignores
    overwritten files anyway: ``cloudFiles.allowOverwrites`` is false). Returns [(target, status)].
    """
    results = []
    for local, target in landing_targets(settings, source_dir, sources):
        if not Path(local).is_file():
            raise FileNotFoundError(local)
        if exists(target) and not overwrite:
            results.append((target, "skipped (exists)"))
            continue
        copy(local, target)
        results.append((target, "copied"))
    return results
