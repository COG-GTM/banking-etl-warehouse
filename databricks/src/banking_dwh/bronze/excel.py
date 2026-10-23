"""Excel transaction file -> Bronze ``file_transaction_excel`` (Ticket 5).

Ports the ``tFileInputExcel_1`` branch of Talend ``Load_FactTransaction``: one sheet
(``Sheet1``), one header row, columns from the first column on, ``DIE_ON_ERROR=false``
(bad rows go to a reject flow instead of failing the job). The Talend schema types are
``transaction_id/account_id/amount/branch_id`` Integer, ``transaction_type`` String and
``transaction_date`` Date, with ``dd-MM-yyyy HH:mm:ss`` for text cells.

Two readers produce the same all-STRING frame, then share one typing step:

* ``openpyxl`` (default, portable): reads the workbook with openpyxl on the driver from a
  local or ``/Volumes`` FUSE path. Works in local tests, on serverless and on classic.
  The file is tiny (7 rows); this is a control-file reader, not a big-data one.
* ``native``: Databricks' built-in Excel data source (``spark.read.format("excel")``,
  DBR 17.x / serverless), read with ``inferColumnTypes=false`` because it does not rescue
  type mismatches. ``com.crealytics.spark.excel`` is the classic-cluster alternative
  (Maven library), not installable on serverless.

Typing never fails the load: a cell that can't be parsed becomes NULL in its column and
its raw text is kept in ``_rescued_data`` (JSON object keyed by column), as are columns
that aren't in the contract. Loads are idempotent per file: rows for the files in the
batch replace any rows previously ingested from the same ``_source_file``.
"""

from __future__ import annotations

import os
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any, Iterable, Sequence

from ..common.config import EnvConfig
from ..contracts.ddl import create_table_sql
from ..contracts.tables import BRONZE_FILE_TRANSACTION_EXCEL

CONTRACT = BRONZE_FILE_TRANSACTION_EXCEL
SOURCE_COLUMNS = tuple(CONTRACT.business_columns)
INT_COLUMNS = ("transaction_id", "account_id", "branch_id")
EXCEL_EXTENSIONS = (".xlsx", ".xlsm")
DEFAULT_SHEET = "Sheet1"
RAW_FILE_COLUMN = "_source_file"
_RAW_EXTRA = "_extra_columns"

# Text formats accepted for transaction_date, in priority order. The first is the Talend
# column pattern; the M/d/yy forms are what Excel shows for the m/d/yy h:mm cell format
# (the native reader returns that display text when reading dates as STRING, and renders
# midnight / single-digit hours without a leading zero, e.g. "2024-01-03 0:00:00").
TIMESTAMP_FORMATS = (
    "dd-MM-yyyy HH:mm:ss",
    "dd-MM-yyyy H:mm:ss",
    "yyyy-MM-dd HH:mm:ss",
    "yyyy-MM-dd H:mm:ss",
    "yyyy-MM-dd'T'HH:mm:ss",
    "dd-MM-yyyy HH:mm",
    "dd-MM-yyyy H:mm",
    "yyyy-MM-dd HH:mm",
    "yyyy-MM-dd H:mm",
    "M/d/yy H:mm",
    "M/d/yyyy H:mm",
    "M/d/yy H:mm:ss",
    "M/d/yyyy H:mm:ss",
    "yyyy-MM-dd",
    "dd-MM-yyyy",
)
# Excel serial dates (days since 1899-12-30), only when a date cell lost its format.
_DECIMAL_19_4_MAX = Decimal("999999999999999.9999")
_EXCEL_EPOCH_OFFSET_DAYS = 25569  # 1970-01-01 as an Excel serial
_EXCEL_SERIAL_MIN, _EXCEL_SERIAL_MAX = 1, 2958465  # 1900-01-01 .. 9999-12-31


# --------------------------------------------------------------------------------------
# File discovery
# --------------------------------------------------------------------------------------
def normalize_path(path: str) -> str:
    """``dbfs:/Volumes/..`` and ``file:/..`` -> plain path, so both readers agree on lineage."""
    for prefix in ("dbfs:", "file://", "file:"):
        if path.startswith(prefix):
            path = path[len(prefix):]
            break
    return path


def list_excel_files(source: str) -> list[str]:
    """A single workbook path, or every .xlsx/.xlsm in a folder (sorted, non-recursive)."""
    source = normalize_path(source)
    if os.path.isfile(source):
        return [source]
    if not os.path.isdir(source):
        raise FileNotFoundError(f"Excel landing path not found: {source}")
    return sorted(
        os.path.join(source, f)
        for f in os.listdir(source)
        if f.lower().endswith(EXCEL_EXTENSIONS) and not f.startswith(("~$", "."))
    )


# --------------------------------------------------------------------------------------
# Reader 1: openpyxl -> all-STRING rows
# --------------------------------------------------------------------------------------
def cell_to_text(value: Any) -> str | None:
    """Canonical text for a cell value, so the Spark typing step sees one format per type."""
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(Decimal(repr(value)), "f")
    if isinstance(value, (int, Decimal)):
        return str(value)
    text = str(value)
    return text if text.strip() != "" else None


def _header_key(value: Any) -> str:
    return str(value).strip().lower() if value is not None else ""


def read_workbook_rows(path: str, sheet: str = DEFAULT_SHEET, header_rows: int = 1) -> list[dict[str, str | None]]:
    """Rows of ``sheet`` as {header: text}. Header names are trimmed and lower-cased.

    Fully empty rows are skipped. Unknown header names are kept (they end up in
    ``_rescued_data``); duplicate or blank headers get positional names (``_c<N>``).
    """
    import openpyxl

    wb = openpyxl.load_workbook(normalize_path(path), read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"Sheet {sheet!r} not found in {path}; sheets: {wb.sheetnames}")
        rows = wb[sheet].iter_rows(values_only=True)
        header: list[str] = []
        for _ in range(header_rows):
            header = [_header_key(v) for v in next(rows, ())]
        seen: set[str] = set()
        names = []
        for i, h in enumerate(header):
            name = h if h and h not in seen else f"_c{i}"
            seen.add(name)
            names.append(name)
        out = []
        for values in rows:
            texts = [cell_to_text(v) for v in values]
            if all(t is None for t in texts):
                continue
            while len(names) < len(texts):
                names.append(f"_c{len(names)}")
            out.append({names[i]: t for i, t in enumerate(texts)})
        return out
    finally:
        wb.close()


def _string_frame(spark, records: Sequence[dict[str, Any]]):
    """All-STRING frame: contract columns, ``_source_file`` and a JSON map of extra columns."""
    import json

    from pyspark.sql.types import StringType, StructField, StructType

    schema = StructType(
        [StructField(c, StringType(), True) for c in SOURCE_COLUMNS]
        + [StructField(RAW_FILE_COLUMN, StringType(), False), StructField(_RAW_EXTRA, StringType(), True)]
    )
    data = []
    for r in records:
        extra = {k: v for k, v in r.items() if k not in SOURCE_COLUMNS and k != RAW_FILE_COLUMN and v is not None}
        data.append(
            tuple(r.get(c) for c in SOURCE_COLUMNS)
            + (r[RAW_FILE_COLUMN], json.dumps(extra, sort_keys=True) if extra else None)
        )
    return spark.createDataFrame(data, schema)


def read_openpyxl(spark, files: Iterable[str], sheet: str = DEFAULT_SHEET, header_rows: int = 1):
    records = []
    for f in files:
        src = normalize_path(f)
        for row in read_workbook_rows(src, sheet, header_rows):
            records.append({**row, RAW_FILE_COLUMN: src})
    return _string_frame(spark, records)


# --------------------------------------------------------------------------------------
# Reader 2: Databricks native Excel source -> all-STRING frame
# --------------------------------------------------------------------------------------
def native_reader_options(sheet: str = DEFAULT_SHEET, header_rows: int = 1) -> dict[str, str]:
    """Options for ``spark.read.format("excel")``. ``dataAddress`` is the bare sheet name
    (``Sheet1!A1`` fails with EXCEL_DATA_ADDRESS_SHEET_NOT_FOUND on serverless)."""
    return {"headerRows": str(header_rows), "dataAddress": sheet, "inferColumnTypes": "false"}


def read_native(spark, files: Sequence[str], sheet: str = DEFAULT_SHEET, header_rows: int = 1):
    from pyspark.sql import functions as F

    raw = spark.read.format("excel").options(**native_reader_options(sheet, header_rows)).load(list(files))
    path = F.regexp_replace(F.col("_metadata.file_path"), r"^(dbfs:|file://|file:)", "")
    return project_native(raw, path)


def project_native(raw, path):
    """Map native-reader columns onto SOURCE_COLUMNS (+ file path, + JSON of unknown columns).

    A row is kept when any known *or* unknown column has a value, matching the openpyxl reader,
    so a row whose only data is in an extra column still lands with it in ``_rescued_data``.
    """
    from pyspark.sql import functions as F

    by_key = {}
    for name in raw.columns:
        by_key.setdefault(name.strip().lower(), name)
    cols = [
        (F.col(f"`{by_key[c]}`").cast("string") if c in by_key else F.lit(None).cast("string")).alias(c)
        for c in SOURCE_COLUMNS
    ]
    extras = [n for k, n in by_key.items() if k not in SOURCE_COLUMNS]
    extra_json = (
        F.to_json(F.map_filter(
            F.create_map(*[x for n in extras for x in (F.lit(n.strip().lower()), F.col(f"`{n}`").cast("string"))]),
            lambda k, v: v.isNotNull(),
        ))
        if extras else F.lit(None).cast("string")
    )
    non_empty = F.coalesce(*[F.col(c) for c in (*SOURCE_COLUMNS, _RAW_EXTRA)]).isNotNull()
    return (
        raw.select(*cols, path.alias(RAW_FILE_COLUMN), extra_json.alias(_RAW_EXTRA))
        .withColumn(_RAW_EXTRA, F.when(F.col(_RAW_EXTRA) != "{}", F.col(_RAW_EXTRA)))
        .where(non_empty)
    )


# --------------------------------------------------------------------------------------
# Shared typing + rescue (Spark Connect safe: no Column.try_cast, no RDD APIs)
# --------------------------------------------------------------------------------------
def _trimmed(c: str):
    from pyspark.sql import functions as F

    v = F.trim(F.col(c))
    return F.when(v != "", v)


def _int_expr(c: str):
    """INT only for whole numbers in range ("6", "6.0"); "6.5", "x" -> NULL (rescued)."""
    from pyspark.sql import functions as F

    d = F.expr(f"try_cast(trim(`{c}`) AS DECIMAL(38,10))")
    ok = d.isNotNull() & (d == F.floor(d)) & d.between(-2147483648, 2147483647)
    return F.when(ok, d.cast("int"))


def _decimal_expr(c: str):
    from pyspark.sql import functions as F

    rounded = F.round(F.expr(f"try_cast(trim(`{c}`) AS DECIMAL(38,10))"), 4)
    # Range-check after rounding: 999999999999999.99995 rounds past DECIMAL(19,4) and must be rescued.
    ok = rounded.isNotNull() & (F.abs(rounded) <= F.lit(_DECIMAL_19_4_MAX))
    return F.when(ok, rounded.cast("decimal(19,4)"))


def _timestamp_expr(c: str):
    from pyspark.sql import functions as F

    text = F.trim(F.col(c))
    parsed = [F.call_function("try_to_timestamp", text, F.lit(fmt)) for fmt in TIMESTAMP_FORMATS]
    serial = F.expr(f"try_cast(trim(`{c}`) AS DOUBLE)")
    from_serial = F.when(
        F.trim(F.col(c)).rlike(r"^[0-9]+(\.[0-9]+)?$")
        & serial.between(_EXCEL_SERIAL_MIN, _EXCEL_SERIAL_MAX),
        F.timestamp_seconds(F.round((serial - _EXCEL_EPOCH_OFFSET_DAYS) * 86400).cast("long")),
    )
    return F.coalesce(*parsed, from_serial)


def type_and_rescue(raw):
    """All-STRING frame (contract columns + ``_source_file`` + extras JSON) -> contract-typed.

    Returns every contract column except ``_ingested_at``; failed casts are NULL and
    recorded in ``_rescued_data`` together with any extra columns.
    """
    from pyspark.sql import functions as F

    typed = {}
    for c in SOURCE_COLUMNS:
        if c in INT_COLUMNS:
            typed[c] = _int_expr(c)
        elif c == "amount":
            typed[c] = _decimal_expr(c)
        elif c == "transaction_date":
            typed[c] = _timestamp_expr(c)
        else:
            typed[c] = F.col(c).cast("string")

    stage = raw.select(
        *[F.col(c).alias(f"__raw_{c}") for c in SOURCE_COLUMNS],
        *[typed[c].alias(c) for c in SOURCE_COLUMNS],
        F.col(RAW_FILE_COLUMN),
        F.col(_RAW_EXTRA),
    )
    failed = F.create_map(*[
        x for c in SOURCE_COLUMNS
        for x in (F.lit(c), F.when(_trimmed(f"__raw_{c}").isNotNull() & F.col(c).isNull(), F.col(f"__raw_{c}")))
    ])
    extra = F.from_json(F.col(_RAW_EXTRA), "map<string,string>")
    merged = F.map_filter(
        F.map_concat(failed, F.coalesce(extra, F.create_map().cast("map<string,string>"))),
        lambda k, v: v.isNotNull(),
    )
    rescued = F.when(F.size(merged) > 0, F.to_json(merged))
    return stage.select(
        *[F.col(c) for c in SOURCE_COLUMNS],
        rescued.alias("_rescued_data"),
        F.col(RAW_FILE_COLUMN),
    )


def with_lineage(typed, ingested_at=None):
    """Add ``_ingested_at`` and project to the contract's column order and types."""
    from pyspark.sql import functions as F

    ts = F.current_timestamp() if ingested_at is None else F.lit(ingested_at).cast("timestamp")
    df = typed.withColumn("_ingested_at", ts)
    return df.select(*[F.col(c.name).cast(c.type).alias(c.name) for c in CONTRACT.columns])


# --------------------------------------------------------------------------------------
# Load
# --------------------------------------------------------------------------------------
def _sql_str(value: str) -> str:
    return "'" + value.replace("\\", "\\\\").replace("'", "\\'") + "'"


def replace_predicate(files: Sequence[str]) -> str:
    return f"{RAW_FILE_COLUMN} IN ({', '.join(_sql_str(normalize_path(f)) for f in files)})"


def ensure_table(spark, cfg: EnvConfig, *, unity_catalog: bool = True) -> str:
    spark.sql(create_table_sql(cfg, CONTRACT, unity_catalog=unity_catalog))
    return cfg.table(CONTRACT.layer, CONTRACT.name)


def build_bronze_frame(spark, files: Sequence[str], *, reader: str = "openpyxl",
                       sheet: str = DEFAULT_SHEET, header_rows: int = 1, ingested_at=None):
    if reader == "openpyxl":
        raw = read_openpyxl(spark, files, sheet, header_rows)
    elif reader == "native":
        raw = read_native(spark, files, sheet, header_rows)
    else:
        raise ValueError(f"Unknown reader {reader!r}; expected 'openpyxl' or 'native'")
    return with_lineage(type_and_rescue(raw), ingested_at)


def ingest_excel(spark, cfg: EnvConfig, source: str | None = None, *, reader: str = "openpyxl",
                 sheet: str = DEFAULT_SHEET, header_rows: int = 1, unity_catalog: bool = True,
                 ingested_at=None) -> dict[str, Any]:
    """Load every workbook under ``source`` (default: the excel_transaction landing folder).

    Rows from those files replace earlier rows with the same ``_source_file`` in one Delta
    commit (``replaceWhere``), so a rerun of the same file never duplicates rows.
    """
    from pyspark.sql import functions as F

    source = source or cfg.landing_path("excel_transaction")
    files = [normalize_path(f) for f in list_excel_files(source)]
    table = ensure_table(spark, cfg, unity_catalog=unity_catalog)
    summary: dict[str, Any] = {"table": table, "source": source, "reader": reader, "sheet": sheet,
                               "files": files, "rows_written": 0, "rescued_rows": 0}
    if not files:
        return summary

    df = build_bronze_frame(spark, files, reader=reader, sheet=sheet, header_rows=header_rows,
                            ingested_at=ingested_at)
    (df.write.format("delta").mode("overwrite")
       .option("replaceWhere", replace_predicate(files))
       .saveAsTable(table))

    written = spark.table(table).where(F.col(RAW_FILE_COLUMN).isin(files))
    stats = written.agg(
        F.count(F.lit(1)).alias("rows"),
        F.sum(F.when(F.col("_rescued_data").isNotNull(), 1).otherwise(0)).alias("rescued"),
    ).collect()[0]
    summary.update(rows_written=int(stats["rows"]), rescued_rows=int(stats["rescued"] or 0),
                   table_rows=spark.table(table).count())
    return summary
